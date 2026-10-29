"""
review — ChatGPT-svaret tillbaka in i arket under AppTest (PR D).

Inklistrat CSV-svar eller uppladdad fil → kontroll → skillnadstabell med "Använd" (förkryssad för high/medium med
källa, utan varning och utan att skriva över arkets värde) → "Skriv in N ändringar" skriver i sessionens rad, kortets
widgets ritas om med de nya värdena och ingenting sparas till GitHub (storage.save_json får inte anropas). Källan
'ChatGPT: …' följer värdet, inte fältet, och Tiggres katalysatorer skrivs utan tomma rader. Durrett-arket får
Datapoints med källa 'ChatGPT: …' och källtyp 'unsupported' när den saknas, Scorecardet läser bara in kortets egna
fält, och Thors Arken-flik är bara läsning. Inget nätverk: lagringen, sifferuppdateringen och Gisten är fejkade.
"""
import copy
import json
import os
import sys

import pytest
import streamlit

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import storage  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STORES = {
    "producers": {"producers": [{"id": "p1", "ticker": "BOL.ST", "name": "Boliden", "commodity": "Guld",
                                 "date": "2026-09-01", "price": 3000.0}],
                  "royalty": [{"id": "r1", "ticker": "FNV", "name": "Franco", "level": 1, "date": "2026-09-01"}]},
    "scoring": {"sprott": [], "durrett": []},
    "insider": {"signals": []},
    "tiggre": {"candidates": [{"id": "c1", "ticker": "LOB", "name": "Lobo", "added": "2026-09-01", "screen": {},
                               "downside": -40.0, "factors": {}, "catalysts": []}],
               "positions": [], "closed": [], "parked": []},
    "holdings": {"swing": [], "ovtlyr": [], "long": [], "momentum": [], "tiggre": [], "closed": [], "cash": 0},
    "scorecard": {"cards": {}},
    "confidence": {"companies": {}, "commodity_overrides": {}},
}
HEAD = "Ticker;Fältnyckel;Värde;Enhet;Källa;Källtyp;Datum;Säkerhet;Kommentar"
RULE_ANSWER = "\n".join([
    "Här är svaret:", "```csv", HEAD,
    "BOL.ST;rule:unit_cost;1 250,5;USD/oz;Q2-rapporten s. 12;primary;2026-08-01;high;AISC 2026",
    "BOL.ST;rule:mine_life;14;år;43-101 2025;primary;2025-03-01;medium;",
    "BOL.ST;rule:kapitaldisciplin;Ja;;Årsredovisning 2025;primary;2026-02-01;high;Återköp pågår",
    "BOL.ST;rule:nav_musd;2100;MUSD;Analytiker;secondary;2026-07-01;low;osäker",
    "BOL.ST;rule:position_pct;4;%;;;;;",
    "NEM;rule:ev_musd;5;MUSD;x;primary;2026-01-01;high;",
    "```"])


def _same(a, b) -> bool:
    return json.dumps(a, sort_keys=True, default=str) == json.dumps(b, sort_keys=True, default=str)


@pytest.fixture
def page(monkeypatch):
    """Arken ur en fejkad lagring — och 💾 får aldrig anropas av importen."""
    stores = copy.deepcopy(STORES)
    monkeypatch.setattr(storage, "session_load", lambda name, default=None, legacy_file=None:
                        streamlit.session_state.setdefault(name, copy.deepcopy(stores.get(name, default))))
    monkeypatch.setattr(storage, "load_error", lambda name: None)
    monkeypatch.setattr(storage, "is_dirty", lambda name: False)
    monkeypatch.setattr(storage, "last_saved", lambda name: None)

    def _no_save(*a, **k):
        raise AssertionError("importen får aldrig spara till GitHub")
    monkeypatch.setattr(storage, "save_json", _no_save)
    monkeypatch.setattr(storage, "save_session", _no_save)
    import gist_storage
    import refresh_ui
    import screens_ui
    monkeypatch.setattr(refresh_ui, "load_refresh", lambda: {})
    monkeypatch.setattr(screens_ui, "load_screens", lambda: {})
    monkeypatch.setattr(gist_storage, "load_blob", lambda *a, **k: None)
    monkeypatch.setenv("RT_TEST_ROOT", ROOT)

    def run(module: str, fn: str, kw: dict, state: dict = None):
        from streamlit.testing.v1 import AppTest
        monkeypatch.setenv("RT_MOD", module)
        monkeypatch.setenv("RT_FN", fn)
        monkeypatch.setenv("RT_KW", json.dumps(kw))

        def app():
            import importlib
            import json as _j
            import os as _o
            import sys as _s
            _s.path.insert(0, _o.environ["RT_TEST_ROOT"])
            getattr(importlib.import_module(_o.environ["RT_MOD"]), _o.environ["RT_FN"])(**_j.loads(_o.environ["RT_KW"]))

        at = AppTest.from_function(app, default_timeout=90)
        for k, v in (state or {}).items():
            at.session_state[k] = v
        at.run()
        assert not at.exception, (module, at.exception)
        return at
    return run


def _rule(at) -> dict:
    return at.session_state["producers"]["producers"][0]


def test_pasted_answer_shows_the_diff_and_writes_only_after_the_button(page):
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    at.text_area(key="rv_imp_txt_rule_p1_0").set_value(RULE_ANSWER).run()
    assert not at.exception, at.exception
    assert _rule(at).get("unit_cost") is None                          # inget skrivs innan knappen
    warns = " ".join(w.value for w in at.warning)
    assert "annan ticker (NEM)" in warns
    caps = " ".join(c.value for c in at.caption)
    assert "Bara du — beslutas av dig" in caps                         # positionen rörs aldrig
    go = at.button(key="rv_imp_go_rule_p1")
    assert go.label == "Skriv in 3 ändringar"                          # nav_musd har säkerhet low → ej förkryssad
    go.click().run()
    assert not at.exception, at.exception
    row = _rule(at)
    assert row["unit_cost"] == 1250.5 and row["mine_life"] == 14.0 and row["kapitaldisciplin"] is True
    assert row.get("nav_musd") is None and row.get("position_pct") is None
    assert row["sources"]["unit_cost"]["source"] == "ChatGPT: Q2-rapporten s. 12 · 2026-08-01"
    assert at.number_input(key="pr_c_p1").value == 1250.5               # widgeten ritades om med nya värdet
    assert at.checkbox(key="pr_ch_p1_kapitaldisciplin").value is True
    assert any("3 ändringar inskrivna i arket" in s.value for s in at.success)
    assert at.text_area(key="rv_imp_txt_rule_p1_1").value == ""         # importen är tömd
    import producers
    assert producers.producer_score(row) is not None                    # grindarna räknar på det inlästa som förut


def test_unticking_a_row_leaves_it_out(page):
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    at.text_area(key="rv_imp_txt_rule_p1_0").set_value(RULE_ANSWER).run()
    tab = next(k for k in at.session_state.filtered_state if str(k).startswith("rv_imp_tab_rule_p1_0_"))
    at.session_state[tab] = {"edited_rows": {0: {"Använd": False}, 3: {"Använd": True}}, "added_rows": [],
                             "deleted_rows": []}
    at.run()
    assert at.button(key="rv_imp_go_rule_p1").label == "Skriv in 3 ändringar"
    at.button(key="rv_imp_go_rule_p1").click().run()
    row = _rule(at)
    assert row.get("unit_cost") is None and row["nav_musd"] == 2100.0     # du valde själv


def test_unit_mismatch_warns_and_is_not_ticked(page):
    """Prototypens fälla: AISC i USD/lb mot guldpriset i USD/oz — varning, ingen omräkning, inget kryss."""
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    answer = "\n".join([HEAD, "BOL.ST;rule:unit_cost;4,1;USD/lb;Rapport;primary;2026-08-01;high;"])
    at.text_area(key="rv_imp_txt_rule_p1_0").set_value(answer).run()
    go = at.button(key="rv_imp_go_rule_p1")
    assert go.disabled and go.label == "Skriv in 0 ändringar"
    table = at.dataframe[0].value                                       # skillnadstabellen
    assert list(table["Använd"]) == [False] and "lb mot oz" in table["Varning"][0]
    assert table["Nytt"][0] == "4,1"                                     # talet står kvar — inget räknat om


def test_an_answer_that_overwrites_your_value_is_unticked_and_flagged(page):
    """Råvarupriset står redan i arket (3000): säkerhet high med källa kryssar inte i bytet — du gör det själv.
    Tabellen skriver talen på svenska och visar fältnyckeln under sitt eget namn."""
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    answer = "\n".join([HEAD, "BOL.ST;rule:price;3100,5;USD/oz;LBMA 2026-09-30;primary;2026-09-30;high;",
                        "BOL.ST;rule:unit_cost;1250,5;USD/oz;Q2-rapporten;primary;2026-08-01;high;"])
    at.text_area(key="rv_imp_txt_rule_p1_0").set_value(answer).run()
    assert not at.exception, at.exception
    table = at.dataframe[0].value
    assert "Ark" not in table.columns and list(table["Fältnyckel"]) == ["rule:price", "rule:unit_cost"]
    rows = {k: i for i, k in enumerate(table["Fältnyckel"])}
    price, cost = rows["rule:price"], rows["rule:unit_cost"]
    assert not table["Använd"][price]
    assert "skriver över värdet i arket" in table["Varning"][price]
    assert table["Nu"][price] == "3 000" and table["Nytt"][price] == "3 100,5"
    assert bool(table["Använd"][cost]) and table["Varning"][cost] == "" and table["Nytt"][cost] == "1 250,5"
    go = at.button(key="rv_imp_go_rule_p1")
    assert go.label == "Skriv in 1 ändring"
    go.click().run()
    assert _rule(at)["price"] == 3000.0 and _rule(at)["unit_cost"] == 1250.5          # ditt pris står kvar


def test_the_chatgpt_source_is_not_kept_on_a_value_you_change(page):
    """Källan 'ChatGPT: …' gäller det inlästa talet. Skriver du ett eget i kortet följer den inte med till nästa
    underlag ('Källa nu')."""
    from review import export as ex
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    at.text_area(key="rv_imp_txt_rule_p1_0").set_value(RULE_ANSWER).run()
    at.button(key="rv_imp_go_rule_p1").click().run()

    def _source():
        row = _rule(at)
        e = ex.build("BOL.ST", {"producers": {"producers": [row], "royalty": []}}, ["rule"], rows={"rule": row})
        return {r["key"]: r["source"] for r in e.rows}["rule:unit_cost"]
    assert _source() == "ChatGPT: Q2-rapporten s. 12 · 2026-08-01"
    at.number_input(key="pr_c_p1").set_value(1300.0).run()
    assert not at.exception, at.exception
    assert _rule(at)["unit_cost"] == 1300.0 and _source() == ""


def test_uploaded_csv_file_is_read_like_the_pasted_answer(page):
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    data = "\n".join([HEAD, "BOL.ST;rule:mine_life;9;år;43-101;primary;2025-03-01;high;"]).encode("utf-8-sig")
    at.file_uploader(key="rv_imp_file_rule_p1_0").set_value(("svar.csv", data, "text/csv")).run()
    assert not at.exception, at.exception
    at.button(key="rv_imp_go_rule_p1").click().run()
    assert _rule(at)["mine_life"] == 9.0
    assert at.number_input(key="pr_life_p1").value == 9.0


def test_uploaded_xlsx_round_trip(page):
    pytest.importorskip("openpyxl")
    import io
    from openpyxl import load_workbook
    from review import export as ex
    row = copy.deepcopy(STORES["producers"]["producers"][0])
    e = ex.build("BOL.ST", {"producers": {"producers": [row], "royalty": []}}, ["rule"], rows={"rule": row})
    wb = load_workbook(io.BytesIO(ex.to_xlsx(e.rows, "prompt")))
    ws = wb[ex.SHEET_FILL]
    head = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(head)}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col["Fältnyckel"]).value == "rule:rp_ratio":
            ws.cell(r, col["Fyll i: Värde"]).value = 11
            ws.cell(r, col["Fyll i: Enhet"]).value = "år"
            ws.cell(r, col["Fyll i: Källa (URL/dokument+sida)"]).value = "Reserve report"
            ws.cell(r, col["Fyll i: Säkerhet (high/medium/low)"]).value = "high"
    buf = io.BytesIO()
    wb.save(buf)
    at = page("producers", "render_producers_page", {"sheet": "Rick Rule"})
    at.file_uploader(key="rv_imp_file_rule_p1_0").set_value(("svar.xlsx", buf.getvalue(), ex.XLSX_MIME)).run()
    assert not at.exception, at.exception
    at.button(key="rv_imp_go_rule_p1").click().run()
    assert _rule(at)["rp_ratio"] == 11.0


def test_tiggre_catalysts_from_the_answer_get_id_and_status(page):
    at = page("tiggre", "render_tiggre_page", {})
    answer = "\n".join([HEAD, "LOB;tiggre:catalysts.0.name;Finansieringsbesked;;PR 2026-09;primary;2026-09-01;high;",
                        "LOB;tiggre:catalysts.0.date;2027-Q1;;PR 2026-09;primary;2026-09-01;high;",
                        "LOB;tiggre:screen.fs;Ja;;DFS 2025;primary;2025-11-01;high;",
                        "LOB;tiggre:factors.finansiering;1;;PR;primary;2026-09-01;medium;Delvis"])
    at.text_area(key="rv_imp_txt_tiggre_c1_0").set_value(answer).run()
    at.button(key="rv_imp_go_tiggre_c1").click().run()
    assert not at.exception, at.exception
    cand = at.session_state["tiggre"]["candidates"][0]
    cat = cand["catalysts"][0]
    assert cat["name"] == "Finansieringsbesked" and cat["date"] == "2027-Q1" and cat["status"] == "Väntar" and cat["id"]
    assert cand["screen"]["fs"] is True and cand["factors"]["finansiering"] == 1
    assert at.checkbox(key="tg_sc_c1_fs").value is True


def test_a_catalyst_answer_into_an_empty_list_leaves_one_row_and_nothing_is_written_on_render(page):
    """Granskningens fall: bara tiggre:catalysts.1.date i en tom lista. Förr blev plats 1 en tom rad utan id som
    ritades '? / ?' och fick status när arket ritades — nu en rad med id och status, och en ny ritning ändrar
    ingenting."""
    import tiggre
    at = page("tiggre", "render_tiggre_page", {})
    answer = "\n".join([HEAD, "LOB;tiggre:catalysts.1.date;2027-03;;PR 2026-09;primary;2026-09-01;high;"])
    at.text_area(key="rv_imp_txt_tiggre_c1_0").set_value(answer).run()
    at.button(key="rv_imp_go_tiggre_c1").click().run()
    assert not at.exception, at.exception
    (cat,) = at.session_state["tiggre"]["candidates"][0]["catalysts"]
    assert cat["date"] == "2027-03" and cat["id"] and cat["status"] == tiggre.CAT_WAITING
    before = copy.deepcopy(at.session_state["tiggre"])
    at.run()
    assert not at.exception, at.exception
    assert _same(at.session_state["tiggre"], before)                        # att rita skriver inget
    assert at.selectbox(key=f"tg_catst_c1_{cat['id']}").value == tiggre.CAT_WAITING


def test_durrett_sheet_import_writes_datapoints_with_chatgpt_source(monkeypatch):
    import durrett_cases as dcs
    import refresh_ui
    import screens_ui
    import streamlit as st
    from confidence import store as cs
    from streamlit.testing.v1 import AppTest

    data = cs.default()
    company = dcs.gold_producer()
    cs.put(data, company)
    monkeypatch.setattr(storage, "session_load", lambda name, default=None, legacy_file=None:
                        st.session_state.setdefault(name, {"confidence": data}.get(name, default)))
    monkeypatch.setattr(storage, "load_error", lambda name: None)
    monkeypatch.setattr(storage, "is_dirty", lambda name: False)
    monkeypatch.setattr(storage, "last_saved", lambda name: None)
    monkeypatch.setattr(refresh_ui, "load_refresh", lambda: {})
    monkeypatch.setattr(screens_ui, "load_screens", lambda: {})
    monkeypatch.setenv("RT_TEST_ROOT", ROOT)

    def app():
        import os as _o
        import sys as _s
        _s.path.insert(0, _o.environ["RT_TEST_ROOT"])
        from engines.durrett.ui import render_durrett_page
        render_durrett_page()

    at = AppTest.from_function(app, default_timeout=90)
    at.run()
    assert not at.exception, at.exception
    t = company.ticker
    answer = "\n".join([HEAD, f"{t};confidence:fields.aisc;1 350;USD/oz;MD&A Q2 2026 s. 4;;2026-08-10;high;",
                        f"{t};confidence:fields.debt_maturity_year;2029;år;Årsredovisning 2025 s. 61;primary;2026-03-01;"
                        f"medium;", f"{t};confidence:fields.irr_pct;31;%;FS 2024;primary;2024-05-01;medium;"])
    at.text_area(key=f"rv_imp_txt_durrett_{t}_0").set_value(answer).run()
    assert not at.exception, at.exception
    # AISC står redan i arket — bytet kryssas inte i av sig självt, du kryssar det
    table = at.dataframe[0].value
    i = list(table["Fältnyckel"]).index("confidence:fields.aisc")
    assert not table["Använd"][i] and "skriver över värdet i arket" in table["Varning"][i]
    tab = next(k for k in at.session_state.filtered_state if str(k).startswith(f"rv_imp_tab_durrett_{t}_0_"))
    at.session_state[tab] = {"edited_rows": {i: {"Använd": True}}, "added_rows": [], "deleted_rows": []}
    at.run()
    at.button(key=f"rv_imp_go_durrett_{t}").click().run()
    assert not at.exception, at.exception
    f = at.session_state["confidence"]["companies"][t]["fields"]
    assert f["aisc"]["value"] == 1350.0 and f["aisc"]["source"] == "ChatGPT: MD&A Q2 2026 s. 4"
    assert f["aisc"]["source_type"] == "unsupported" and f["aisc"]["kind"] == "ESTIMATE"   # källtyp saknades
    p = f["debt_maturity_year"]
    assert p["value"] == 2029 and p["source_type"] == "primary" and p["pub_date"] == "2026-03-01"
    assert p["confidence"] == "medium" and "irr_pct" not in f            # IRR hör till utvecklare — visas inte här
    assert at.text_input(key=f"ds_{t}_aisc").value == "1350"                  # Arket ritades om


def test_scorecard_reads_only_its_own_fields_and_creates_the_card_then(page):
    at = page("scorecard", "render_scorecard_page", {})
    key = "boliden|BOL.ST"
    assert at.session_state["scorecard"]["cards"] == {}                        # orört kort skrivs inte
    answer = "\n".join([HEAD, "BOL.ST;scorecard:valuation;P/NAV 0,4;;Analys 2026;secondary;2026-09-01;high;",
                        "BOL.ST;rule:mine_life;14;år;43-101;primary;2025-03-01;high;",
                        "BOL.ST;scorecard:sell_rule;Stopp −20 %;;;;;high;"])
    at.text_area(key=f"rv_imp_txt_scorecard_{key}_0").set_value(answer).run()
    caps = " ".join(c.value for c in at.caption)
    assert "hör till Rick Rule — läs in svaret där" in caps and "Bara du" in caps
    at.button(key=f"rv_imp_go_scorecard_{key}").click().run()
    assert not at.exception, at.exception
    card = at.session_state["scorecard"]["cards"][key]
    assert card["valuation"] == "P/NAV 0,4" and card["strategy"] == "producenter" and "sell_rule" not in card
    assert at.session_state["producers"]["producers"][0].get("mine_life") is None
    assert at.text_input(key=f"sc_val_{key}").value == "P/NAV 0,4"


def test_thor_arken_tab_has_no_import(monkeypatch):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    import test_review_chatgpt_ui as base
    monkeypatch.setattr(storage, "session_load", lambda name, default=None, legacy_file=None:
                        st.session_state.setdefault(name, copy.deepcopy(base.THOR_STORES.get(name, default))))
    at = AppTest.from_function(base._thor_script, default_timeout=120)
    at.session_state["_stores"], at.session_state["_strategy"] = copy.deepcopy(base.THOR_STORES), "rule"
    at.run()
    assert not at.exception, at.exception
    assert not at.get("file_uploader") and not [t for t in at.text_area if str(t.key).startswith("rv_imp_")]
    assert any("Svaret läses tillbaka i arket" in c.value for c in at.caption)

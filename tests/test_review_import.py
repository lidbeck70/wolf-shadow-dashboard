"""
review/answers + review/units + review/importer — ChatGPT-svaret läses, kontrolleras och skrivs in (rena funktioner).

Läsningen: kort CSV-svar (';', ',' eller tabb, i kodstaket, med BOM), hela exporten med 'Fyll i'-kolumnerna, JSON
(rader eller extraktionens schema) och Excel. Tolkningen: decimalkomma, tusental, %, Ja/Nej, val, 0–2, heltal,
datum och månad — och "vet inte" (null, N/A, –) är tomt. Kontrollen: annan ticker, okänd nyckel och ogiltigt värde
avvisas, 'Bara du' och andra ark hoppas över, en enhet/valuta/skala som skiljer sig varnar (aldrig omräkning),
Lukacs target-yield prövas mot FCF-klassens band, och bara ändrade värden blir ändringar. Inskrivningen: platta ark
får värdet och 'sources', Durrett/Wolf Asymmetry ett Datapoint ('ChatGPT: …', källtyp 'unsupported' när den saknas).
"""
import copy
import io
import json
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import controls as ctl  # noqa: E402
import producers as prod  # noqa: E402
from review import answers as an  # noqa: E402
from review import catalog as cat  # noqa: E402
from review import export as ex  # noqa: E402
from review import importer as im  # noqa: E402
from review import units  # noqa: E402

HEAD = ex.ANSWER_HEADER


def _rule_row(**kw) -> dict:
    return {"id": "p1", "ticker": "BOL.ST", "name": "Boliden", "commodity": "Koppar", "price": 4.1, **kw}


def _exp(sheet="rule", row=None, ticker="BOL.ST"):
    row = row if row is not None else _rule_row()
    stores = {"producers": {"producers": [row], "royalty": []}}
    return ex.build(ticker, stores, [sheet], rows={sheet: row}), row


def _check(lines, sheet="rule", row=None, writable=None):
    e, r = _exp(sheet, row)
    return im.check(an.read_text("\n".join([HEAD] + list(lines))), e, writable), r, e


# ── Läsningen ────────────────────────────────────────────────────────────────
def test_short_answer_in_every_delimiter_fence_and_bom():
    rows = ["BOL.ST;rule:unit_cost;2,5;USD/lb;Q2 s. 3;primary;2026-08-01;high;AISC"]
    comma = 'BOL.ST,rule:unit_cost,"2,5",USD/lb,Q2 s. 3,primary,2026-08-01,high,AISC'
    for text in ("\n".join([HEAD] + rows),
                 "Svar:\n```csv\n" + "\n".join([HEAD] + rows) + "\n```\nLycka till",
                 "\N{ZERO WIDTH NO-BREAK SPACE}" + "\n".join([HEAD] + rows),
                 "\n".join([HEAD.replace(";", "\t")] + [r.replace(";", "\t") for r in rows]),
                 "\n".join([HEAD.replace(";", ","), comma])):
        got = an.read_text(text)
        assert len(got) == 1, text
        a = got[0]
        assert (a.ticker, a.key, a.value, a.unit, a.source, a.source_type, a.date, a.confidence, a.comment) == (
            "BOL.ST", "rule:unit_cost", "2,5", "USD/lb", "Q2 s. 3", "primary", "2026-08-01", "high", "AISC")
        assert a.line == 2
    english = "Ticker;Key;Value;Unit;Source;Source type;Date;Confidence;Comment\nBOL.ST;rule:mine_life;9;år;x;;;;"
    assert an.read_text(english)[0].value == "9"
    with pytest.raises(an.ImportFormatError):
        an.read_text("hej\nbara text")


def test_the_full_export_reads_only_the_fill_columns():
    e, _r = _exp()
    rows = copy.deepcopy(e.rows)
    for r in rows:
        if r["key"] == "rule:mine_life":
            r.update(in_value="12", in_unit="år", in_source="43-101", in_confidence="high")
    got = an.read_file(ex.to_csv(rows), "svar.csv")
    hit = [a for a in got if a.key == "rule:mine_life"]
    assert len(hit) == 1 and hit[0].value == "12" and hit[0].unit == "år" and hit[0].source == "43-101"
    # arkets egna kolumner (Värde nu, Enhet) blandas aldrig ihop med svaret
    price = [a for a in got if a.key == "rule:price"]
    assert all(a.value == "" for a in price)
    assert len(got) == len(rows)


def test_json_rows_and_the_extraction_schema():
    rows = [{"Ticker": "BOL.ST", "Fältnyckel": "rule:mine_life", "Värde": 7, "Källa": "x", "Säkerhet": "high"}]
    a = an.read_text(json.dumps(rows))[0]
    assert a.key == "rule:mine_life" and a.value == 7 and a.confidence == "high"
    assert an.read_text(json.dumps({"rows": rows}))[0].value == 7
    schema = {"fields": {"aisc": {"value": 1250, "unit": "USD/oz", "page": 4, "quote": "AISC $1,250",
                                  "confidence": "high"}, "npv_musd": {"value": None}}}
    got = an.read_text("```json\n" + json.dumps(schema) + "\n```")
    assert len(got) == 1 and got[0].key == "aisc" and got[0].source == "s. 4" and got[0].comment == "AISC $1,250"
    with pytest.raises(an.ImportFormatError):
        an.read_text("[1, 2")
    with pytest.raises(an.ImportFormatError):
        an.read_text(json.dumps({"x": 1}))


def test_xlsx_round_trip_reads_the_fill_sheet():
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    e, _r = _exp()
    wb = load_workbook(io.BytesIO(ex.to_xlsx(e.rows, "prompt")))
    ws = wb[ex.SHEET_FILL]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col["Fältnyckel"]).value == "rule:unit_cost":
            ws.cell(r, col["Fyll i: Värde"]).value = 2.6                         # Excel-cellen är ett tal
            ws.cell(r, col["Fyll i: Datum (ÅÅÅÅ-MM-DD)"]).value = __import__("datetime").date(2026, 8, 1)
    buf = io.BytesIO()
    wb.save(buf)
    got = {a.key: a for a in an.read_file(buf.getvalue(), "svar.xlsx")}
    assert got["rule:unit_cost"].value == 2.6 and got["rule:unit_cost"].date == "2026-08-01"


def test_without_openpyxl_an_excel_file_is_a_clear_error(monkeypatch):
    import builtins
    real = builtins.__import__

    def fake(name, *a, **k):
        if name.startswith("openpyxl"):
            raise ImportError("nej")
        return real(name, *a, **k)
    monkeypatch.setattr(builtins, "__import__", fake)
    with pytest.raises(an.ImportFormatError, match="openpyxl"):
        an.read_file(b"PK\x03\x04rest", "svar.xlsx")


# ── Tolkningen ───────────────────────────────────────────────────────────────
@pytest.mark.parametrize("raw,value,unit,warns", [
    ("1 250,5", 1250.5, "", False), ("1,250.5", 1250.5, "", False), ("1.250,5", 1250.5, "", False),
    ("4,1", 4.1, "", False), ("0,350", 0.35, "", False), ("12 345 678", 12345678.0, "", False),
    ("1.250.000", 1250000.0, "", False), ("-3,2", -3.2, "", False), ("−3,2", -3.2, "", False),
    ("45 %", 45.0, "%", False), ("1250 USD/oz", 1250.0, "USD/oz", False), ("3.8x", 3.8, "x", False),
    ("1,250", 1.25, "", True), ("~1200", 1200.0, "", True), (1250, 1250.0, "", False),
    ("1\xa0250,5", 1250.5, "", False), ("1\N{NARROW NO-BREAK SPACE}250,5", 1250.5, "", False),
])
def test_numbers_read_swedish_and_english_without_guessing(raw, value, unit, warns):
    v, u, w = an.parse_number(raw)
    assert v == pytest.approx(value) and u == unit and bool(w) is warns


@pytest.mark.parametrize("raw", ["abc", "12-24", "ca tolv", True])
def test_text_is_not_a_number(raw):
    v, _u, err = an.parse_number(raw)
    assert v is None and "inte ett tal" in err


def _field(sheet, path, row=None):
    return next(f for f in cat.fields_for(sheet, row or {}, position_pct=99.0) if f.path == path)


def test_coerce_every_kind():
    ds = _field("rule", ctl.DS_FIELDS[0].key)
    assert an.coerce(ds, "2")[0] == 2 and an.coerce(ds, "1 — 12–24 mån")[0] == 1 and an.coerce(ds, 0.0)[0] == 0
    assert "0, 1 eller 2" in an.coerce(ds, "3")[1]
    jur = _field("rule", "jurisdiktion")
    assert an.coerce(jur, "Ja")[0] is True and an.coerce(jur, "nej")[0] is False and an.coerce(jur, 1)[0] is True
    assert "Ja eller Nej" in an.coerce(jur, "kanske")[1]
    life = _field("rule", "mine_life")
    assert an.coerce(life, "12 år") == (12.0, "", "år", "")
    assert "under 0" in an.coerce(life, "-1")[1]
    role = _field("insider", "role")
    assert an.coerce(role, role.choices[0].lower())[0] == role.choices[0]
    assert "ogiltigt val" in an.coerce(role, "Kassör")[1]
    level = _field("royalty", "level")
    assert an.coerce(level, "2")[0] == 2 and an.coerce(level, 3.0)[0] == 3
    n = _field("insider", "insiders")
    assert an.coerce(n, "3")[0] == 3 and "heltal" in an.coerce(n, "2,5")[1]
    month = _field("tiggre", "catalysts.0.date", {"catalysts": []})
    assert an.coerce(month, "2027-03")[0] == "2027-03" and an.coerce(month, "q1 2027")[0] == "2027-Q1"
    assert an.coerce(month, "2027 H2")[0] == "2027-H2" and "ÅÅÅÅ-MM" in an.coerce(month, "snart")[1]
    text = _field("rule", "nav_source")
    assert an.coerce(text, "  NI 43-101 2025, 5 %  ")[0] == "NI 43-101 2025, 5 %"
    for blank in ("", "  ", "null", "N/A", "–", None, float("nan")):
        assert an.coerce(life, blank) == (None, "", "", ""), blank
    assert an.iso_date("2026/08/01") == "2026-08-01" and an.iso_date("2026-13-01") is None


# ── Enheterna ────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("text,parts", [
    ("USD/oz", ("usd", "oz", "")), ("US$/oz", ("usd", "oz", "")), ("USD/lb U3O8", ("usd", "lb", "")),
    ("MCAD", ("cad", "", "M")), ("C$ M", ("cad", "", "M")), ("USD bn", ("usd", "", "B")),
    ("MUSD/år", ("usd", "år", "M")), ("mkr", ("sek", "", "M")), ("Moz AuEq", ("", "", "")), ("", ("", "", "")),
])
def test_unit_parts(text, parts):
    assert units.parts(text) == parts


def test_unit_warnings_never_convert():
    f = _field("rule", "unit_cost")
    copper = {"commodity": "Koppar"}
    assert "lb" in units.commodity_unit(copper)
    assert any("oz mot lb" in w for w in units.warnings(f, "USD/oz", copper))
    assert units.warnings(f, "USD/lb", copper) == []
    assert any("Enhet saknas" in w for w in units.warnings(f, "", copper))
    assert any("Valutan CAD" in w for w in units.warnings(f, "CAD/lb", copper))
    # okänd råvara: jämför med råvarupriset i samma svar
    assert any("oz mot lb" in w for w in units.warnings(f, "USD/oz", {}, {"rule:price": "USD/lb"}))
    nav = _field("rule", "nav_musd")
    assert any("Valutan CAD" in w for w in units.warnings(nav, "MCAD", {}))
    assert any("miljarder" in w for w in units.warnings(nav, "USD bn", {}))
    assert units.warnings(nav, "MUSD", {}) == [] and units.warnings(nav, "", {}) == []
    life = _field("rule", "mine_life")
    assert units.warnings(life, "years", {}) == [] and units.warnings(life, "%", {})


# ── Kontrollen ───────────────────────────────────────────────────────────────
def test_check_rejects_other_tickers_unknown_keys_and_bad_values():
    res, _r, _e = _check(["NEM;rule:mine_life;9;år;x;primary;;high;",
                          "BOL.ST;rule:hittepa;9;;x;;;high;",
                          "BOL.ST;mine_life_typo;9;;x;;;high;",
                          "BOL.ST;rule:jurisdiktion;kanske;;x;;;high;",
                          ";rule:mine_life;9;år;43-101;primary;2025-03-01;high;"])
    texts = [t for _l, _k, t in res.issues]
    assert any("annan ticker (NEM)" in t for t in texts) and texts.count("okänd fältnyckel") == 2
    assert any("Ja eller Nej" in t for t in texts)
    assert [c.key for c in res.changes] == ["rule:mine_life"]                   # tom ticker godtas


def test_bara_du_other_sheets_and_hidden_controls_are_skipped_with_a_reason():
    res, _r, _e = _check(["BOL.ST;rule:position_pct;4;%;x;;;high;",
                          "BOL.ST;tiggre:nav;5;MUSD;x;;;high;",
                          f"BOL.ST;rule:{ctl.AQS_FIELDS[0].key};2;;x;;;high;"])
    why = {k: t for _l, k, t in res.ignored}
    assert "Bara du" in why["rule:position_pct"] and "hör till Tiggre" in why["tiggre:nav"]
    assert "visas inte i det här kortet" in why[f"rule:{ctl.AQS_FIELDS[0].key}"]       # AQS under 2 %
    assert not res.changes and not res.issues


def test_only_changed_values_become_changes_and_defaults_follow_confidence_and_source():
    row = _rule_row(mine_life=12.0)
    res, _r, _e = _check(["BOL.ST;rule:mine_life;12;år;x;primary;;high;",
                          "BOL.ST;rule:rp_ratio;9;år;Rapport;primary;2026-01-01;high;",
                          "BOL.ST;rule:ev_musd;900;MUSD;Rapport;primary;2026-01-01;low;",
                          "BOL.ST;rule:nav_musd;1100;MUSD;;;;high;",
                          "BOL.ST;rule:insyn;Ja;;Proxy;primary;2026-xx;medium;"], row=row)
    assert res.same == 1
    use = {c.key: c.default_use for c in res.changes}
    assert use == {"rule:rp_ratio": True, "rule:ev_musd": False, "rule:nav_musd": False, "rule:insyn": False}
    notes = {c.key: c.warnings for c in res.changes}
    assert "källa saknas" in notes["rule:nav_musd"] and any("ÅÅÅÅ-MM-DD" in w for w in notes["rule:insyn"])


def test_bare_keys_and_extraction_aliases_resolve_without_conversion():
    res, _r, _e = _check(["BOL.ST;mine_life;14;år;x;primary;;high;",
                          "BOL.ST;aisc;2,4;USD/lb;x;primary;;high;",              # AISC = unit_cost (faktor 1)
                          "BOL.ST;quarterly_burn_musd;5;MUSD;x;;;high;"])           # faktor 4 — räknas inte om
    got = {c.key: c.new for c in res.changes}
    assert got == {"rule:mine_life": 14.0, "rule:unit_cost": 2.4}
    assert any(t == "okänd fältnyckel" for _l, k, t in res.issues if k == "quarterly_burn_musd")


def test_lukacs_target_yield_follows_the_fcf_class_in_the_same_answer():
    row = _rule_row(position_pct=5.0, fcf_kvalitet="A")
    e, _r = _exp(row=row)
    band_a = cat.fields_for("rule", row)
    assert any(f.path == "fv.Base.target_yield" for f in band_a)
    import lukacs
    lo_c, hi_c = lukacs.yield_band("C")
    outside_a = hi_c                                                         # utanför A:s band, inom C:s
    ans = an.read_text("\n".join([HEAD, f"BOL.ST;rule:fv.Base.target_yield;{outside_a:g};%;Modell;;;high;"]))
    res = im.check(ans, e)
    assert not res.changes and "utanför klass A" in res.issues[0][2]        # fel, inte varning (lukacs)
    ans = an.read_text("\n".join([HEAD, "BOL.ST;rule:fcf_kvalitet;C;;Modell;;;high;",
                                  f"BOL.ST;rule:fv.Base.target_yield;{outside_a:g};%;Modell;;;high;"]))
    res = im.check(ans, e)
    assert {c.key for c in res.changes} == {"rule:fcf_kvalitet", "rule:fv.Base.target_yield"}


def test_a_sheet_without_the_company_cannot_be_written():
    stores = {"producers": {"producers": [], "royalty": []}}
    e = ex.build("BOL.ST", stores, ["rule"])
    res = im.check(an.read_text("\n".join([HEAD, "BOL.ST;rule:mine_life;9;år;x;;;high;"])), e)
    assert not res.changes and "finns inte i Rick Rule" in res.issues[0][2]


# ── Inskrivningen ────────────────────────────────────────────────────────────
def test_apply_flat_rows_get_values_and_sources_and_the_gates_read_them():
    res, row, _e = _check(["BOL.ST;rule:unit_cost;2,1;USD/lb;Q2 s. 12;primary;2026-08-01;high;",
                           "BOL.ST;rule:jurisdiktion;Ja;;Fraser 2025;secondary;2025-12-01;high;",
                           "BOL.ST;rule:kapitaldisciplin;Ja;;Årsredovisning;primary;;high;",
                           "BOL.ST;rule:mine_life;14;år;43-101;primary;2025-03-01;high;",
                           f"BOL.ST;rule:{ctl.DS_FIELDS[0].key};1;;Kassa/burn;primary;2026-08-01;medium;"])
    assert im.apply(res.changes, {"rule": row}) == 5
    assert row["unit_cost"] == 2.1 and row["jurisdiktion"] is True and row[ctl.DS_FIELDS[0].key] == 1
    assert row["sources"]["unit_cost"] == "ChatGPT: Q2 s. 12 · 2026-08-01"
    assert row["sources"]["kapitaldisciplin"] == "ChatGPT: Årsredovisning"
    score = prod.producer_score(row)                                          # arkets egen regel, oförändrad
    assert score == prod.margin_points(prod.margin_pct(4.1, 2.1)) + 2
    assert prod.producer_verdict(score, prod.asset_dying(row)).label == prod.P_BUY
    # exporten visar nu källan för det inlästa
    e2 = ex.build("BOL.ST", {"producers": {"producers": [row], "royalty": []}}, ["rule"], rows={"rule": row})
    src = {r["key"]: r["source"] for r in e2.rows}
    assert src["rule:unit_cost"] == "ChatGPT: Q2 s. 12 · 2026-08-01"


def test_the_copper_trap_changes_nothing_unless_you_tick_it():
    """AISC i USD/oz mot kopparpris i USD/lb: varnas, inte förkryssad — och arket står kvar."""
    res, row, _e = _check(["BOL.ST;rule:unit_cost;1250,5;USD/oz;Rapport;primary;2026-08-01;high;"])
    (ch,) = res.changes
    assert not ch.default_use and any("oz mot lb" in w for w in ch.warnings) and ch.new == 1250.5
    im.apply([c for c in res.changes if c.default_use], {"rule": row})
    assert row.get("unit_cost") is None


def test_apply_confidence_writes_datapoints_with_source_and_unsupported_default():
    import durrett_cases as dcs
    from confidence import store as cs
    data = cs.default()
    company = dcs.gold_producer()
    company.fields.pop("debt_maturity_year", None)
    cs.put(data, company)
    t = company.ticker
    live = data["companies"][t]
    e = ex.build(t, {"confidence": data}, ["confidence"], rows={"confidence": copy.deepcopy(live)})
    ans = an.read_text("\n".join([HEAD, f"{t};confidence:fields.debt_maturity_year;2029;;Årsredovisning s. 61;;"
                                        f"2026-03-01;high;Lånet förfaller", f"{t};fields.dilution_score;3;;Egen bedömning;"
                                        f"weak;;medium;"]))
    res = im.check(ans, e)
    assert im.apply(res.changes, {"confidence": live}) == len(res.changes) == 2
    p = live["fields"]["debt_maturity_year"]
    assert p["value"] == 2029 and p["source"] == "ChatGPT: Årsredovisning s. 61" and p["source_type"] == "unsupported"
    assert p["kind"] == "ESTIMATE" and p["pub_date"] == "2026-03-01" and p["note"] == "Lånet förfaller"
    d = live["fields"]["dilution_score"]
    assert d["kind"] == "MODELLED" and d["source_type"] == "weak"              # en bedömning lagras som modellerad
    # samma svar läst i Wolf Asymmetry-kortet (samma lager) hittar fältet där
    e2 = ex.build(t, {"confidence": data}, ["asymmetry"], rows={"asymmetry": copy.deepcopy(live)})
    f, why = im.resolve("confidence:fields.aisc", e2.fields)
    assert f is not None and f.sheet == "asymmetry", why


def test_widget_keys_follow_the_row_id_and_prefixes_but_never_the_import_itself():
    keys = ["pr_c_p1", "csm_p1_Bear_price", "pr_c_p12", "rv_imp_txt_rule_p1_0", "sc_val_boliden|BOL.ST",
            "cf_GPR_aisc_v", "ds_GPR_aisc", "durrett_row_sub_GPR", "arken_focus"]
    assert im.widget_keys(keys, ["p1"]) == ["pr_c_p1", "csm_p1_Bear_price"]
    assert im.widget_keys(keys, ["boliden|BOL.ST"]) == ["sc_val_boliden|BOL.ST"]
    assert im.widget_keys(keys, (), ("cf_GPR_", "ds_GPR_")) == ["cf_GPR_aisc_v", "ds_GPR_aisc"]


def test_known_path_tells_hidden_from_made_up():
    assert im.known_path("rule", ctl.AQS_FIELDS[0].key) and im.known_path("rule", "csm.Bull.fcf_musd")
    assert im.known_path("tiggre", "catalysts.4.date") and im.known_path("confidence", "fields.aisc")
    assert not im.known_path("rule", "hittepa") and not im.known_path("confidence", "fields.hittepa")

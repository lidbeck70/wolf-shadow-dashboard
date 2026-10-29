"""
review/answers.py — läs ChatGPT-svaret och tolka varje värde i fältets typ.

  read_file(data, name)   det ifyllda arket: Excel (.xlsx, fliken 'Fyll i') eller CSV/TXT/JSON
  read_text(text)         inklistrat svar: CSV-kodblocket (';', ',' eller tabb) eller JSON — även extraktionens schema
  coerce(field, raw)      ett värde i fältets typ: decimalkomma, tusentalsmellanrum, %, Ja/Nej, val, 0–2, datum

Rubrikraden känns igen både i hela exporten (kolumnerna 'Fyll i: …') och i det korta svaret
'Ticker;Fältnyckel;Värde;…'. Rena funktioner, ingen Streamlit.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Optional

from review import catalog as cat

MAX_TEXT = 1000
_BOM = "\N{ZERO WIDTH NO-BREAK SPACE}"


class ImportFormatError(ValueError):
    """Svaret gick inte att läsa som arket (ingen rubrikrad, trasig JSON, Excel utan openpyxl)."""


@dataclass
class Answer:
    """En ifylld rad ur svaret — råa strängar (eller Excel-cellens tal)."""
    line: int
    ticker: str = ""
    key: str = ""
    value: Any = ""
    unit: str = ""
    source: str = ""
    source_type: str = ""
    date: str = ""
    confidence: str = ""
    comment: str = ""


# ── Rubrikraden ──────────────────────────────────────────────────────────────
_ANSWER_COLS = {
    "värde": "value", "value": "value", "enhet": "unit", "unit": "unit",
    "källa": "source", "kalla": "source", "source": "source",
    "källtyp": "source_type", "kalltyp": "source_type", "source type": "source_type", "source_type": "source_type",
    "sourcetype": "source_type", "datum": "date", "date": "date",
    "säkerhet": "confidence", "sakerhet": "confidence", "confidence": "confidence",
    "kommentar": "comment", "kommentar/citat": "comment", "comment": "comment", "citat": "comment", "quote": "comment",
}
_META_COLS = {"ticker": "ticker", "fältnyckel": "key", "faltnyckel": "key", "field key": "key", "fieldkey": "key",
              "key": "key", "nyckel": "key"}
_FILL = "fyll i:"
_HEAD = re.compile(r"(?i)f[äa]ltnyckel|\bkey\b")


def _norm(h) -> str:
    s = re.sub(r"\s+", " ", str(h or "").replace(_BOM, "").strip().lower())
    return re.sub(r"\s*\(.*\)\s*$", "", s).strip()


def _columns(header: list) -> dict:
    """{kolumnindex: fältnamn}. Hela exporten (kolumnerna 'Fyll i: …') eller det korta svaret (Ticker;Fältnyckel;…)."""
    names = [_norm(h) for h in header]
    full = any(n.startswith(_FILL) for n in names)
    out: dict = {}
    for i, n in enumerate(names):
        key = None
        if n.startswith(_FILL):
            key = _ANSWER_COLS.get(n[len(_FILL):].strip())
        elif n in _META_COLS:
            key = _META_COLS[n]
        elif not full:                      # i hela exporten är 'Enhet' och 'Värde nu' arkets, inte svaret
            key = _ANSWER_COLS.get(n)
        if key and key not in out.values():
            out[i] = key
    return out


def cell(v) -> Any:
    """En cell som tal (Excel) eller text; datum som ÅÅÅÅ-MM-DD."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if v is None:
        return ""
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v).strip()


def _answers(rows: list) -> list:
    """Raderna under rubrikraden som Answer (rubrikraden hittas på 'Fältnyckel'/'Key')."""
    head_at, cols_at = None, {}
    for i, r in enumerate(rows):
        cols = _columns(list(r))
        if "key" in cols.values():
            head_at, cols_at = i, cols
            break
    if head_at is None:
        raise ImportFormatError("Hittar ingen rubrikrad med Fältnyckel — svara med arket eller CSV-blocket med "
                                "rubrikraden ur prompten.")
    out = []
    for j, r in enumerate(rows[head_at + 1:], head_at + 2):          # radnummer som i filen (rubriken = rad 1)
        r = list(r)
        a = Answer(line=j)
        for i, name in cols_at.items():
            if i < len(r):
                setattr(a, name, cell(r[i]) if name == "value" else str(cell(r[i])))
        if not str(a.key).strip() and cat.is_empty(a.value):
            continue
        out.append(a)
    return out


def _from_json(obj) -> list:
    if isinstance(obj, dict) and isinstance(obj.get("fields"), dict):     # extraktionens schema (ai/extract_prompt)
        out = []
        for i, (k, v) in enumerate(obj["fields"].items(), 1):
            if not isinstance(v, dict) or v.get("value") in (None, ""):
                continue
            page = v.get("page")
            out.append(Answer(line=i, ticker=str(obj.get("ticker") or ""), key=str(k), value=cell(v.get("value")),
                              unit=str(v.get("unit") or ""), source=f"s. {page}" if page not in (None, "") else "",
                              confidence=str(v.get("confidence") or ""), comment=str(v.get("quote") or "")))
        return out
    if isinstance(obj, dict):
        obj = obj.get("rows") or obj.get("answers") or obj.get("svar")
    if not isinstance(obj, list) or not all(isinstance(r, dict) for r in obj):
        raise ImportFormatError("JSON-svaret är varken en lista med rader eller extraktionens schema.")
    keys = list(dict.fromkeys(k for r in obj for k in r))
    return _answers([keys] + [[r.get(k, "") for k in keys] for r in obj])


_FENCE = re.compile(r"```[^\n`]*\n(.*?)```", re.S)


def read_text(text: str) -> list:
    """Inklistrat svar: CSV-kodblocket eller JSON (ett kodstaket runt räcker att det finns)."""
    raw = str(text or "").replace(_BOM, "").strip()
    blocks = [b.strip() for b in _FENCE.findall(raw)] or [raw]
    body = next((b for b in blocks if b.startswith(("{", "["))), None)
    if body is None and raw.startswith(("{", "[")):
        body = raw
    if body is not None:
        try:
            obj = json.loads(body)
        except ValueError as exc:
            raise ImportFormatError(f"JSON-svaret gick inte att läsa: {exc}") from None
        return _from_json(obj)
    body = next((b for b in blocks if _HEAD.search(b)), blocks[0])
    lines = body.splitlines()
    head = next((ln for ln in lines if _HEAD.search(ln)), lines[0] if lines else "")
    delim = max((";", "\t", ","), key=lambda d: (head.count(d), d == ";"))
    return _answers(list(csv.reader(io.StringIO(body), delimiter=delim)))


def read_file(data: bytes, name: str = "") -> list:
    """Den uppladdade filen: Excel känns igen på zip-signaturen, allt annat läses som text (UTF-8, annars cp1252)."""
    data = bytes(data or b"")
    if data[:2] == b"PK":
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportFormatError("Excel kräver paketet openpyxl — spara som CSV eller klistra in svaret.") from None
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ImportFormatError(f"Excel-filen gick inte att öppna: {exc}") from None
        from review.export import SHEET_FILL
        ws = wb[SHEET_FILL] if SHEET_FILL in wb.sheetnames else wb.active
        return _answers([list(r) for r in ws.iter_rows(values_only=True)])
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return read_text(data.decode(enc))
        except UnicodeDecodeError:
            continue
    raise ImportFormatError(f"{name or 'Filen'} gick inte att läsa som text.")


# ── Ett värde i fältets typ ──────────────────────────────────────────────────
_SPACES = r"\s\xa0\N{NARROW NO-BREAK SPACE}'"
_NUM = re.compile(rf"^([+\-−–]?)\s*(\d[\d{_SPACES}.,]*)\s*(.*)$")
_APPROX = re.compile(r"^(?:~|≈|ca\.?\s|c:a\s|cirka\s|approx\.?\s|about\s)", re.I)
_UNIT_TEXT = re.compile(r"^[%A-Za-zÅÄÖåäö$€£/×.\s0-9]+$")
_BLANKS = {"", "-", "–", "—", "n/a", "na", "null", "none", "saknas", "okänt", "okant", "?", "tomt"}
_BOOL = {"ja": True, "j": True, "yes": True, "y": True, "true": True, "sant": True, "1": True,
         "nej": False, "n": False, "no": False, "false": False, "falskt": False, "0": False}


def blank(raw) -> bool:
    """Tomt svar: inget värde, eller ett 'vet inte' (null, N/A, –) — ingenting att skriva in."""
    if raw is None or isinstance(raw, bool):
        return raw is None
    if isinstance(raw, (int, float)):
        return raw != raw
    return str(raw).strip().lower() in _BLANKS


def parse_number(raw) -> tuple:
    """(tal, enhet efter talet, varning eller fel) ur '1 250,5', '1,250.5', '45 %', '1250 USD/oz'.

    (None, '', fel) när det inte är ett tal. Ett ensamt skiljetecken med tre siffror efter ('1,250') läses som
    decimal — svaret är svenskt — men varnas, för det kan vara tusental."""
    if isinstance(raw, bool):
        return None, "", "inte ett tal"
    if isinstance(raw, (int, float)):
        return (None, "", "inte ett tal") if raw != raw else (float(raw), "", "")
    s = str(raw or "").strip()
    warn = ""
    if _APPROX.match(s):
        s, warn = _APPROX.sub("", s).strip(), "ungefärligt värde i svaret"
    m = _NUM.match(s)
    if not m:
        return None, "", f"inte ett tal (”{raw}”)"
    sign, rest = m.group(1), m.group(3).strip()
    digits = re.sub(rf"[{_SPACES}]", "", m.group(2)).rstrip(".,")
    if "," in digits and "." in digits:
        dec = "," if digits.rfind(",") > digits.rfind(".") else "."
        digits = digits.replace("." if dec == "," else ",", "").replace(",", ".")
    elif digits.count(",") > 1 or digits.count(".") > 1:
        digits = digits.replace(",", "").replace(".", "")                # tusentalstecken
    elif "," in digits or "." in digits:
        whole, frac = digits.split("," if "," in digits else ".")
        if len(frac) == 3 and whole not in ("", "0") and len(whole) <= 3:
            warn = f"”{raw}” läst som {whole},{frac} (decimal) — var det tusental? Kontrollera."
        digits = f"{whole}.{frac}"
    try:
        v = float(digits)
    except ValueError:
        return None, "", f"inte ett tal (”{raw}”)"
    if rest and not _UNIT_TEXT.match(rest):
        return None, "", f"inte ett tal (”{raw}”)"
    return (-v if sign in ("-", "−", "–") else v), rest, warn


def _month(s: str) -> Optional[str]:
    s = s.strip()
    if re.match(r"^\d{4}-(0[1-9]|1[0-2])(-(0[1-9]|[12]\d|3[01]))?$", s):
        return s
    m = re.match(r"^(\d{4})\s*-?\s*([QqHh])([1-4])$", s) or re.match(r"^([QqHh])([1-4])\s*-?\s*(\d{4})$", s)
    if not m:
        return None
    g = m.groups()
    year, kind, n = (g[0], g[1], g[2]) if g[0].isdigit() else (g[2], g[0], g[1])
    if kind.upper() == "H" and n not in ("1", "2"):
        return None
    return f"{year}-{kind.upper()}{n}"


def iso_date(raw) -> Optional[str]:
    """ÅÅÅÅ-MM-DD (eller ÅÅÅÅ-MM), annars None."""
    s = str(cell(raw)).strip().replace("/", "-")
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        try:
            return date.fromisoformat(s).isoformat()
        except ValueError:
            return None
    return s if re.match(r"^\d{4}-(0[1-9]|1[0-2])$", s) else None


def _choice(f: cat.SheetField, raw) -> tuple:
    s = str(cell(raw)).strip()
    for c in f.choices:
        if str(c).strip().lower() == s.lower():
            return c, ""
    v, _u, _w = parse_number(raw)
    if v is not None:
        for c in f.choices:
            if isinstance(c, (int, float)) and not isinstance(c, bool) and float(c) == v:
                return c, ""
    m = re.match(r"^\s*(\d+)\s*[—–\-:(]", s)                            # '2 — Nettokassa'
    if m:
        for c in f.choices:
            if str(c) == m.group(1):
                return c, ""
    return None, f"ogiltigt val ”{s}” — tillåtna: {', '.join(str(c) for c in f.choices)}"


def _number(f: cat.SheetField, raw) -> tuple:
    v, unit, warn = parse_number(raw)
    if v is None:
        return None, warn, "", ""
    if f.kind == cat.INT:
        if abs(v - round(v)) > 1e-9:
            return None, f"heltal krävs ({v:g})", "", ""
        v = int(round(v))
    # Lukacs target-yield prövas mot FCF-klassens band i importen (klassen kan bytas i samma svar)
    banded = f.path.endswith(".target_yield")
    if f.lo is not None and v < (0.0 if banded else f.lo):
        return None, f"under {(0.0 if banded else f.lo):g} är inte tillåtet ({v:g})", "", ""
    if f.hi is not None and v > f.hi and not banded:
        return None, f"över {f.hi:g} är inte tillåtet ({v:g})", "", ""
    return v, "", unit, warn


def coerce(f: cat.SheetField, raw) -> tuple:
    """(värde, fel, enhet ur värdet, varning). Tomt svar = (None, '', '', '') — inget att skriva."""
    if blank(raw):
        return None, "", "", ""
    kind = f.kind
    if kind == cat.BOOL:
        if isinstance(raw, (int, float)) and not isinstance(raw, bool) and raw in (0, 1):
            return bool(raw), "", "", ""
        b = _BOOL.get(str(raw).strip().lower())
        return (b, "", "", "") if b is not None else (None, f"Ja eller Nej krävs (”{raw}”)", "", "")
    if kind == cat.SCORE02:
        v, _u, _w = parse_number(raw)
        if v is None:
            m = re.match(r"^\s*([0-2])\s*(?:p|poäng)?\s*(?:[—–\-:(].*)?$", str(raw))
            v = float(m.group(1)) if m else None
        if v is None or v not in (0.0, 1.0, 2.0):
            return None, f"0, 1 eller 2 krävs (”{raw}”)", "", ""
        return int(v), "", "", ""
    if kind == cat.CHOICE or (kind == cat.INT and f.choices):
        v, err = _choice(f, raw)
        return (v, "", "", "") if not err else (None, err, "", "")
    if kind in (cat.NUMBER, cat.INT):
        return _number(f, raw)
    if kind == cat.DATE:
        d = iso_date(raw)
        return (d, "", "", "") if d and len(d) == 10 else (None, f"datum ÅÅÅÅ-MM-DD krävs (”{raw}”)", "", "")
    if kind == cat.MONTH:
        d = _month(str(cell(raw)))
        return (d, "", "", "") if d else (None, f"månad ÅÅÅÅ-MM krävs (”{raw}”)", "", "")
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(cell(raw)).strip()[:MAX_TEXT], "", "", ""

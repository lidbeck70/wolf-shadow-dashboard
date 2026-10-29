"""
review/export.py — "Underlag till ChatGPT": aktiens fält som kalkylark och en prompt att klistra in.

  build(ticker, stores, …)  rader + luckor för de ark aktien finns i (eller de ark som skickas in — ett ark där
                            aktien saknas blir en tom mall). Värdena förifylls ur arken och, där arket är tomt, ur
                            sifferuppdateringens Börsdata-förslag. Det som saknas för köpbeslut står först.
  to_xlsx(rows, prompt)     Excel: fliken 'Fyll i' (låsta frågekolumner, blå ifyllnadskolumner, rullgardiner för
                            Ja/Nej, val och 0–2) och fliken 'Instruktion' med prompten. None utan openpyxl.
  to_csv(rows)              samma rader som semikolonseparerad CSV med BOM (csv_export) — reservutgången
  build_prompt(…)           svensk prompt: fyll bara i 'Fyll i'-kolumnerna på ChatGPT- och Bedömningsrader, källa och
                            datum per värde eller tomt, aldrig enhets- eller valutaomräkning, NAV/NPV efter skatt,
                            0–2 med en mening, rör aldrig 'Bara du'
  filename(ticker, ext)     underlag_<ticker>_<datum>.<ext>

Inget skrivs i arken härifrån. Rena funktioner, ingen Streamlit.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Optional

import controls as ctl
import csv_export
import review_link
import scorecard as sc
import scoring as sco
from review import catalog as cat
from review import missing as mis

# (nyckel, rubrik) — ordningen är filens kolumnordning och ändras inte utan att importen följer med
EXPORT_COLUMNS = [
    ("ticker", "Ticker"), ("sheet", "Ark"), ("key", "Fältnyckel"), ("label", "Fält"), ("kind", "Typ"),
    ("unit", "Enhet"), ("allowed", "Tillåtna värden"), ("value", "Värde nu"), ("source", "Källa nu"),
    ("who", "Vem fyller"), ("required", "Krävs för köpbeslut"), ("hint", "Instruktion/definition"),
    ("in_value", "Fyll i: Värde"), ("in_unit", "Fyll i: Enhet"),
    ("in_source", "Fyll i: Källa (URL/dokument+sida)"),
    ("in_source_type", "Fyll i: Källtyp (primary/independent/secondary/mixed/weak/unsupported)"),
    ("in_date", "Fyll i: Datum (ÅÅÅÅ-MM-DD)"), ("in_confidence", "Fyll i: Säkerhet (high/medium/low)"),
    ("in_comment", "Fyll i: Kommentar/citat"),
]
INPUT_KEYS = ("in_value", "in_unit", "in_source", "in_source_type", "in_date", "in_confidence", "in_comment")
ANSWER_HEADER = "Ticker;Fältnyckel;Värde;Enhet;Källa;Källtyp;Datum;Säkerhet;Kommentar"
SOURCE_TYPES = ("primary", "independent", "secondary", "mixed", "weak", "unsupported")
CONFIDENCES = ("high", "medium", "low")
SHEET_FILL, SHEET_HELP = "Fyll i", "Instruktion"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

KIND_LABEL = {cat.NUMBER: "Tal", cat.INT: "Heltal", cat.SCORE02: "Poäng 0–2", cat.BOOL: "Ja/Nej",
              cat.CHOICE: "Val", cat.TEXT: "Text", cat.DATE: "Datum", cat.MONTH: "Månad"}
YES, NO = "Ja", "Nej"


@dataclass
class Export:
    ticker: str
    name: str = ""
    sheets: list = field(default_factory=list)       # arknycklarna i underlaget
    rows: list = field(default_factory=list)
    missing: list = field(default_factory=list)      # [MissingItem]
    absent: list = field(default_factory=list)       # ark där aktien inte finns — raderna är en tom mall

    @property
    def filled(self) -> int:
        """Fält med ett sparat värde i arket (ett Börsdata-förslag räknas inte — det är inte inskrivet)."""
        return sum(1 for r in self.rows if r.get("_stored"))


# ── Var aktien finns ─────────────────────────────────────────────────────────
def _forms(ticker: str) -> list:
    t = str(ticker or "").strip().upper()
    try:
        from thor.ticker import forms
        out = [f for f in forms(t) if f]
    except Exception:
        out = []
    return out or ([t] if t else [])


def _store(stores: dict, sheet: str) -> dict:
    s = (stores or {}).get(cat.SHEETS[sheet].store)
    return s if isinstance(s, dict) else {}


def find(sheet: str, ticker: str, stores: dict) -> Optional[dict]:
    """Aktiens rad i arket (någon av tickerns former), eller None. Scorecardet: kortet."""
    fs = _forms(ticker)
    if sheet in cat.CONFIDENCE_SHEETS:
        comps = _store(stores, sheet).get("companies") or {}
        return next((comps[f] for f in fs if isinstance(comps.get(f), dict)), None)
    if sheet == "scorecard":
        return card_for(ticker, stores)[1]
    for f in fs:
        row = review_link.find_row(sheet, f, stores or {})
        if row is not None:
            return row
    return None


def card_for(ticker: str, stores: dict) -> tuple:
    """(kortnyckel, kortet, kandidatraden) ur scorecardet — kortet som det visas (strategin förvald)."""
    fs = _forms(ticker)
    cards = _store(stores, "scorecard").get("cards") or {}
    entries = sc.entries_for(fs, stores or {})
    for k, entry in entries.items():
        c = cards.get(k)
        return k, sc._shown(c if isinstance(c, dict) else {}, entry), entry
    for k, c in cards.items():
        if isinstance(c, dict) and any(str(k).endswith("|" + f) for f in fs):
            return k, dict(c), None
    return None, None, None


def sheets_with(ticker: str, stores: dict) -> list:
    """Arken aktien finns i, i katalogens ordning. Durrett 10-steg täcker Wolf Asymmetry (samma lager)."""
    return [s for s in cat.SHEETS if s != "asymmetry" and find(s, ticker, stores) is not None]


# ── Värdena ──────────────────────────────────────────────────────────────────
def _show(v):
    if isinstance(v, bool):
        return YES if v else NO
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (dict, list)) or v is None:
        return ""
    return str(v)


def _fmt(v) -> str:
    return f"{v:g}" if isinstance(v, (int, float)) and not isinstance(v, bool) else str(v)


def allowed(f: cat.SheetField) -> str:
    """Tillåtna värden i klartext — samma lista som rullgardinen i Excel."""
    if f.kind == cat.BOOL:
        return f"{YES}, {NO}"
    if f.kind == cat.SCORE02:
        return "0, 1, 2"
    if f.choices:
        return ", ".join(_fmt(c) for c in f.choices)
    if f.kind == cat.DATE:
        return "ÅÅÅÅ-MM-DD"
    if f.kind == cat.MONTH:
        return "ÅÅÅÅ-MM"
    if f.kind in (cat.NUMBER, cat.INT):
        word = "heltal" if f.kind == cat.INT else "tal"
        if f.lo is not None and f.hi is not None:
            return f"{word} {f.lo:g}–{f.hi:g}"
        if f.lo is not None:
            return f"{word} ≥ {f.lo:g}"
        return word
    return "fritext"


def choice_list(f: cat.SheetField) -> tuple:
    """Rullgardinens val i Excel, eller tomt när fältet är fritt."""
    if f.kind == cat.BOOL:
        return (YES, NO)
    if f.kind == cat.SCORE02:
        return ("0", "1", "2")
    return tuple(_fmt(c) for c in f.choices)


def _suggestion(blob: Optional[dict], sheet: str, row: dict, source_field: str):
    """Sifferuppdateringens förslag (värde, datum, enhet) för raden, eller None."""
    if not blob or not source_field or not isinstance(row, dict) or not row.get("id"):
        return None
    try:
        import refresh_ui
        return refresh_ui.suggestion(blob, cat.SHEETS[sheet].refresh, row, source_field,
                                     unit_field="commodity_unit" if source_field == "commodity_price" else None)
    except Exception:
        return None


def _conf_proposals(blob: Optional[dict], row: dict) -> dict:
    """{fältnyckel: Datapoint} ur sifferuppdateringen för Durrett/Wolf Asymmetry-arkets bolag."""
    if not blob or not isinstance(row, dict) or not row.get("ticker"):
        return {}
    try:
        from confidence.data.models import CompanyInput
        from engines.durrett import refresh as dref
        return {k: p for k, p, _cur in dref.proposals(blob, CompanyInput.from_dict(row))}
    except Exception:
        return {}


def _ds_hints(blob: Optional[dict], sheet: str, row: dict) -> dict:
    """Förslag till DS-delfrågorna ur arkets egna förslagsfunktioner — text i instruktionen, aldrig ett värde."""
    out = {}
    g = _suggestion(blob, sheet, row, "shares_growth_3y_pct")
    if g is not None:
        out["ds_aktier_3ar"] = (f"Börsdata {g[1]}: antalet aktier {g[0]:+.0f} % mot för tre år sedan — motsvarar "
                                f"{ctl.ds_shares_suggestion(g[0])} riskpoäng.")
    if sheet == "sprott":
        pts = ctl.ds_runway_suggestion(sco.runway_years(row.get("cash"), row.get("burn")))
        if pts is not None:
            out["ds_runway"] = f"Runwayen i arket motsvarar {pts} riskpoäng."
    return out


def _rows_for(ticker: str, sheet: str, row: Optional[dict], blob: Optional[dict], stage: Optional[str],
              position_pct, entry: Optional[dict], strategies: tuple) -> tuple:
    r = cat._as_row(row) if row is not None else {}
    fields = cat.fields_for(sheet, r, stage, position_pct, strategies)
    items = mis.missing_for(sheet, r, entry=entry, stage=stage, position_pct=position_pct)
    need = {}
    for it in items:
        if it.path:
            need.setdefault(it.path, it.reason)
    props = _conf_proposals(blob, r) if sheet in cat.CONFIDENCE_SHEETS else {}
    ds_hints = _ds_hints(blob, sheet, r) if sheet in cat.FLAT_SHEETS else {}
    label = cat.SHEETS[sheet].label
    out = []
    for f in fields:
        v = cat.get_value(r, f.path)
        stored = not cat.is_empty(v)
        source, unit, who, hint = "", f.unit, f.who, f.hint
        point = cat.get_point(r, f.path)
        if point is not None and not point.missing:
            source = " · ".join(x for x in (point.source, point.pub_date or "") if x)
            unit = point.unit or unit
        if cat.is_empty(v):
            sug = None
            if f.who == cat.BORSDATA and sheet in cat.FLAT_SHEETS:
                s_ = _suggestion(blob, sheet, r, f.refresh)
                if s_ is not None:
                    sug = (s_[0], f"Börsdata {s_[1]} — förslag, inte inskrivet", s_[2])
            elif sheet in cat.CONFIDENCE_SHEETS and f.path.split(".", 1)[1] in props:
                p = props[f.path.split(".", 1)[1]]
                sug = (p.value, f"{p.source} — förslag, inte inskrivet", p.unit)
            if sug is not None:
                v, source = sug[0], sug[1]
                unit = sug[2] or unit
            elif f.who == cat.BORSDATA and blob is not None:
                who = cat.CHATGPT          # Börsdata har inget tal för raden — då behövs en källa
                hint = ("Börsdata har inget tal för raden — fyll i ur senaste rapporten. " + hint).strip()
        extra = ds_hints.get(f.path)
        if extra:
            hint = f"{hint} {extra}".strip()
        out.append({"ticker": ticker, "sheet": label, "key": f.key, "label": f.label,
                    "kind": KIND_LABEL.get(f.kind, f.kind), "unit": unit, "allowed": allowed(f),
                    "value": _show(v), "source": source, "who": who,
                    "required": f"Ja — {need[f.path]}" if f.path in need else "", "hint": hint,
                    **{k: "" for k in INPUT_KEYS},
                    "_kind": f.kind, "_choices": choice_list(f), "_lo": f.lo, "_hi": f.hi, "_path": f.path,
                    "_sheet": sheet, "_stored": stored})
    return out, items


def _rank(r: dict) -> int:
    """Det som saknas för köpbeslut först, sedan tomma fält ChatGPT kan fylla, sedan dina, sist det ifyllda."""
    if r["required"]:
        return 0
    if cat.is_empty(r["value"]):
        return 2 if r["who"] == cat.BARA_DU else 1
    return 3


def build(ticker: str, stores: dict, sheets=None, stage: Optional[str] = None, blob: Optional[dict] = None,
          name: str = "", rows: Optional[dict] = None, entry: Optional[dict] = None) -> Export:
    """Underlaget för aktien. sheets=None tar arken aktien finns i; ett ark i listan där aktien saknas blir en
    tom mall. blob är sifferuppdateringen (sheets_refresh.json) — None när den inte lästs, och då står tomma
    Börsdata-fält kvar som Börsdata (förslaget finns i arket). rows={ark: rad} låser raden (kortet som ritas)
    i stället för att söka på tickern; rows['scorecard'] är kortet och entry dess kandidatrad (scorecard.collect)."""
    t = str(ticker or "").strip().upper()
    stores = stores or {}
    wanted = [s for s in (sheets if sheets is not None else sheets_with(t, stores)) if s in cat.SHEETS]
    wanted = list(dict.fromkeys(wanted))
    _k, card, found = card_for(t, stores)
    entry = entry if entry is not None else found
    if (rows or {}).get("scorecard") is not None:
        card = rows["scorecard"]
    card_pos = (card or {}).get("position_pct_total")
    card_sheet = cat.STRATEGY_SHEET.get((card or {}).get("strategy") or "")
    exp = Export(t, name)
    for sheet in wanted:
        row = (rows or {}).get(sheet)
        if row is None:
            row = find(sheet, t, stores)
        if row is None and sheet != "scorecard":
            exp.absent.append(sheet)
        if row is not None and not exp.name:
            exp.name = str(cat._as_row(row).get("name") or "")
        if sheet == "scorecard":
            strategies = tuple((entry or {}).get("strategies") or ())
            got, items = _rows_for(t, sheet, card or {}, blob, stage, None, entry, strategies)
        else:
            pos = card_pos if (sheet == card_sheet and cat.SHEETS[sheet].strategy) else None
            got, items = _rows_for(t, sheet, row, blob, stage, pos, None, ())
        exp.sheets.append(sheet)
        exp.rows += got
        exp.missing += items
    exp.rows.sort(key=_rank)
    return exp


def build_rows(ticker: str, stores: dict, sheets=None, stage: Optional[str] = None,
               blob: Optional[dict] = None) -> list:
    """Raderna i underlaget (se build), det som saknas för köpbeslut först."""
    return build(ticker, stores, sheets, stage, blob).rows


# ── Filerna ──────────────────────────────────────────────────────────────────
def filename(ticker: str, ext: str, day: Optional[date] = None) -> str:
    safe = re.sub(r"[^A-Za-z0-9_-]+", "-", str(ticker or "aktie").strip()).strip("-") or "aktie"
    return f"underlag_{safe}_{(day or date.today()).isoformat()}.{ext.lstrip('.')}"


def to_csv(rows: list) -> bytes:
    """Semikolonseparerad CSV med BOM — Excel läser å/ä/ö rätt, och ChatGPT läser den som text."""
    return csv_export.rows_to_csv(rows, EXPORT_COLUMNS).encode("utf-8-sig")


def xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


# Färgerna i arket: blått = fyll i, grått = rör inte
_HEAD_Q, _HEAD_IN, _CELL_IN, _CELL_LOCKED, _CELL_NEED = "3A3F44", "1F4E78", "DDEBF7", "EDEDED", "FCE4D6"
_WIDTH = {"ticker": 10, "sheet": 16, "key": 24, "label": 30, "kind": 10, "unit": 14, "allowed": 18, "value": 14,
          "source": 22, "who": 11, "required": 34, "hint": 60, "in_value": 16, "in_unit": 12, "in_source": 34,
          "in_source_type": 16, "in_date": 14, "in_confidence": 12, "in_comment": 44}


def _quoted(values) -> Optional[str]:
    """Rullgardinens lista som Excel-formel, eller None när den inte går att skriva (kommatecken, för lång)."""
    vals = [str(v) for v in values]
    if not vals or any("," in v or '"' in v for v in vals):
        return None
    text = '"' + ",".join(vals) + '"'
    return text if len(text) <= 255 else None


def to_xlsx(rows: list, prompt: str = "") -> Optional[bytes]:
    """Excel-filen: 'Fyll i' och 'Instruktion'. None när openpyxl inte är installerat (då gäller CSV)."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Protection
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_FILL
    keys = [k for k, _h in EXPORT_COLUMNS]
    col = {k: i for i, k in enumerate(keys, 1)}
    fill = {c: PatternFill("solid", fgColor=c) for c in (_HEAD_Q, _HEAD_IN, _CELL_IN, _CELL_LOCKED, _CELL_NEED)}
    for j, (k, header) in enumerate(EXPORT_COLUMNS, 1):
        c = ws.cell(row=1, column=j, value=header)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill[_HEAD_IN if k in INPUT_KEYS else _HEAD_Q]
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = _WIDTH.get(k, 14)
    ws.row_dimensions[1].height = 48
    wrap = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows or [], 2):
        mine = r.get("who") == cat.BARA_DU
        for k in keys:
            v = r.get(k)
            c = ws.cell(row=i, column=col[k], value=None if v == "" else v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"                 # text, aldrig en formel
            c.alignment = wrap
            if k in INPUT_KEYS:
                c.fill = fill[_CELL_LOCKED if mine else _CELL_IN]
                c.protection = Protection(locked=mine)
            elif k == "required" and v:
                c.fill = fill[_CELL_NEED]

    # rullgardiner och talkontroll — en DataValidation per regel, cellerna läggs till
    dvs: dict = {}

    def _dv(kind: str, formula1: str, formula2: Optional[str] = None, operator: Optional[str] = None):
        key = (kind, formula1, formula2, operator)
        if key not in dvs:
            dv = DataValidation(type=kind, formula1=formula1, formula2=formula2, operator=operator, allow_blank=True)
            dv.error = "Värdet är inte tillåtet här — se Tillåtna värden."
            dv.errorTitle = "Ogiltigt värde"
            dv.showErrorMessage = True
            dvs[key] = dv
            ws.add_data_validation(dv)
        return dvs[key]

    src_list, conf_list = _quoted(SOURCE_TYPES), _quoted(CONFIDENCES)
    for i, r in enumerate(rows or [], 2):
        if r.get("who") == cat.BARA_DU:
            continue
        cell = f"{get_column_letter(col['in_value'])}{i}"
        choices = _quoted(r.get("_choices") or ())
        kind, lo, hi = r.get("_kind"), r.get("_lo"), r.get("_hi")
        if choices:
            _dv("list", choices).add(cell)
        elif kind in (cat.NUMBER, cat.INT):
            t = "whole" if kind == cat.INT else "decimal"
            if lo is not None and hi is not None:
                _dv(t, f"{lo:g}", f"{hi:g}", "between").add(cell)
            elif lo is not None:
                _dv(t, f"{lo:g}", None, "greaterThanOrEqual").add(cell)
        _dv("list", src_list).add(f"{get_column_letter(col['in_source_type'])}{i}")
        _dv("list", conf_list).add(f"{get_column_letter(col['in_confidence'])}{i}")

    ws.freeze_panes = "E2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{len(rows) + 1}"
    # skrivskyddet utan lösenord: frågekolumnerna står still, kolumnbredd, filter och sortering går att ändra
    ws.protection.sheet = True
    for opt in ("formatColumns", "formatRows", "formatCells", "autoFilter", "sort"):
        setattr(ws.protection, opt, False)

    help_ws = wb.create_sheet(SHEET_HELP)
    help_ws.column_dimensions["A"].width = 130
    for i, line in enumerate((prompt or "").splitlines() or [""], 1):
        c = help_ws.cell(row=i, column=1, value=line or None)
        if isinstance(line, str) and line.startswith("="):
            c.data_type = "s"
        c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Prompten ─────────────────────────────────────────────────────────────────
RULES = (
    "Fyll BARA i kolumnerna som börjar med \"Fyll i:\", och bara på rader där \"Vem fyller\" är ChatGPT eller "
    "Bedömning. Rader med Börsdata har redan sitt tal.",
    "Rader med \"Bara du\" rör du aldrig — position, köpgrind, säljregel, kvalitetsgrind och trigger beslutar jag "
    "själv.",
    "Varje värde ska ha en källa (URL, eller dokument + sida) och ett datum (ÅÅÅÅ-MM-DD). Hittar du ingen källa: "
    "lämna raden tom. Gissa aldrig.",
    "Konvertera aldrig enheter eller valutor. Skriv talet i källans enhet och ange enheten i \"Fyll i: Enhet\" — "
    "jag räknar om själv. Står källan i en annan enhet än kolumnen Enhet: säg det i kommentaren.",
    "NAV och NPV ska alltid vara EFTER skatt. Ange studie, diskonteringsränta och metallpris i kommentaren.",
    "Poäng 0–2: välj enligt kriterierna i \"Instruktion/definition\" och motivera med en mening i "
    "\"Fyll i: Kommentar/citat\". Bedömningsrader är förslag som jag granskar.",
    "Ja/Nej: svara Ja eller Nej och citera det som stödjer svaret.",
    "Källtyp: primary (årsredovisning, kvartalsrapport, 43-101/JORC, bolagets egen rapport), independent "
    "(oberoende QP eller analys med egen modell), secondary (välrenommerad sekundärkälla), mixed, weak, "
    "unsupported. Säkerhet: high (står rakt ut), medium (härlett ur två tal), low (otydligt år eller enhet).",
    "Ändra aldrig kolumnerna Ticker, Ark, Fältnyckel, Fält, Värde nu eller Vem fyller, och lägg inte till eller "
    "ta bort rader.",
    "Svara med det ifyllda arket (.xlsx eller .csv). Kan du inte skicka en fil: svara med ett CSV-kodblock med "
    "semikolon och exakt den här rubrikraden, en rad per ifyllt fält:",
)


def build_prompt(ticker: str, name: str, rows: list, strategy: str = "", dom_text: str = "",
                 missing: Optional[list] = None, file_name: str = "") -> str:
    """Prompten att klistra in i ChatGPT tillsammans med filen — fungerar också utan fil (fältlistan står med)."""
    who = {r.get("key"): r.get("who") for r in rows or []}
    sheets = list(dict.fromkeys(r.get("sheet") for r in rows or [] if r.get("sheet")))
    head = f"Underlag för köpbeslut: {name} ({ticker})" if name else f"Underlag för köpbeslut: {ticker}"
    lines = [head + (f" · strategi {strategy}" if strategy else "")]
    if sheets:
        lines.append("Ark: " + ", ".join(sheets))
    lines.append(f"Bifogat: {file_name or filename(ticker, 'xlsx')} — fliken \"{SHEET_FILL}\" "
                 f"(instruktionen står också i fliken \"{SHEET_HELP}\").")
    lines += ["", "Du hjälper mig att fylla i det som saknas i min granskning av aktien. Panelens regler räknar "
                  "poäng, grindar och beslut själva — du hittar fakta och föreslår bedömningar, och jag granskar "
                  "varje rad innan något skrivs in.", "", "Så här fyller du i:"]
    lines += [f"{i}. {rule}" for i, rule in enumerate(RULES, 1)]
    lines.append(ANSWER_HEADER)

    gaps = list(missing or [])
    if gaps:
        lines += ["", "Saknas för köpbeslut (börja här):"]
        for it in gaps:
            sheet = cat.SHEETS[it.sheet].label if it.sheet in cat.SHEETS else it.sheet
            key = f"{it.sheet}:{it.path}" if it.path else ""
            mine = (not it.path) or who.get(key) == cat.BARA_DU
            tail = " [beslutas av mig — fyll inte i]" if mine else f" [{key}]"
            lines.append(f"- {sheet} · {it.label} — {it.reason}{tail}")
    if dom_text:
        lines += ["", "Arkens dom just nu (Thor) — bara bakgrund, ändra inget efter den:", dom_text.strip()]

    todo = [r for r in rows or [] if r.get("who") in (cat.CHATGPT, cat.BEDOMNING) and cat.is_empty(r.get("value"))]
    if todo:
        lines += ["", f"Fält att fylla i ({len(todo)} st — Fältnyckel · fält · enhet · tillåtet):"]
        for r in todo:
            unit = f" · {r['unit']}" if r.get("unit") else ""
            lines.append(f"- {r['key']} · {r['label']}{unit} · {r['allowed']}")
    return "\n".join(lines)
